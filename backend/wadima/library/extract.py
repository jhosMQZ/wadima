"""Extracción de texto, tablas y estructura por formato, y conversión a Markdown estructurado.

Cada extractor devuelve una `Extraction` con el Markdown, el texto plano original (para comparar),
las tablas estructuradas y avisos. Los documentos son datos no confiables: nada de su contenido se
interpreta como instrucción.
"""

from __future__ import annotations

import contextlib
import csv
import io
import json
import re
import shutil
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime, time
from typing import Any

from wadima.library.validate import CODE_EXTS

MAX_TABLE_ROWS_MD = 200
MAX_TABLE_ROWS_JSON = 5000
MAX_COLS = 60
MAX_FORMULAS_MD = 150


class ExtractionError(Exception):
    pass


@dataclass
class Extraction:
    markdown: str
    plain_text: str
    tables: list[dict] = field(default_factory=list)
    pages: int | None = None
    warnings: list[str] = field(default_factory=list)
    needs_ocr: bool = False
    meta: dict = field(default_factory=dict)
    noise_removed: list[str] = field(default_factory=list)


# --------------------------------------------------------------------------- utilidades
def decode_text(content: bytes) -> tuple[str, str]:
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return content.decode(enc), enc
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace"), "utf-8 (con reemplazos)"


def cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        return f"{value:.10g}"
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value)


def md_escape(text: str) -> str:
    return text.replace("|", "\\|").replace("\r", " ").replace("\n", " ").strip()


def md_table(headers: list[str], rows: list[list[str]]) -> str:
    width = max([len(headers)] + [len(r) for r in rows]) if (headers or rows) else 0
    if width == 0:
        return ""
    hdr = [md_escape(h) or f"Col{i + 1}" for i, h in enumerate(headers + [""] * (width - len(headers)))]
    lines = ["| " + " | ".join(hdr) + " |", "|" + "|".join(["---"] * width) + "|"]
    for r in rows:
        r = r + [""] * (width - len(r))
        lines.append("| " + " | ".join(md_escape(c) for c in r) + " |")
    return "\n".join(lines)


def split_header(rows: list[list[str]]) -> tuple[list[str], list[list[str]]]:
    """Detecta la fila de encabezado: primera fila con ≥2 celdas no vacías y mayoría de texto no numérico."""
    for i, r in enumerate(rows[:10]):
        filled = [c for c in r if c.strip()]
        if len(filled) >= 2 and sum(not _is_number(c) for c in filled) >= len(filled) * 0.6:
            return r, rows[i + 1 :]
    return [], rows


def _is_number(text: str) -> bool:
    try:
        float(text.replace(",", ".").replace("%", "").replace("€", "").replace("$", "").strip())
        return True
    except ValueError:
        return False


# --------------------------------------------------------------------------- PDF
def tesseract_available() -> bool:
    return shutil.which("tesseract") is not None


def extract_pdf(content: bytes) -> Extraction:
    import pymupdf
    import pymupdf4llm

    try:
        doc = pymupdf.open(stream=content, filetype="pdf")
    except Exception as exc:  # noqa: BLE001 - PyMuPDF lanza varios tipos
        raise ExtractionError(f"No se pudo abrir el PDF: {exc}") from exc
    if doc.needs_pass:
        raise ExtractionError("El PDF está protegido con contraseña")
    warnings: list[str] = []
    pymupdf4llm.use_layout(False)  # ruta ligera: menos memoria y CPU
    try:
        chunks = pymupdf4llm.to_markdown(doc, page_chunks=True, show_progress=False)
    except Exception as exc:  # noqa: BLE001
        warnings.append(f"Conversión estructurada fallida ({type(exc).__name__}); se usa texto plano")
        chunks = [{"text": page.get_text(), "metadata": {"page": i + 1}} for i, page in enumerate(doc)]
    page_texts = [page.get_text() for page in doc]
    scanned = [
        i for i, page in enumerate(doc) if len(page_texts[i].strip()) < 25 and len(page.get_images(full=False)) > 0
    ]
    needs_ocr = False
    if scanned:
        if tesseract_available():
            for i in scanned:
                try:
                    tp = doc[i].get_textpage_ocr(language="spa+eng", dpi=200, full=True)
                    text = doc[i].get_text(textpage=tp)
                    page_texts[i] = text
                    chunks[i]["text"] = text
                except Exception as exc:  # noqa: BLE001
                    warnings.append(f"OCR fallido en la página {i + 1}: {exc}")
                    needs_ocr = True
            warnings.append(f"Se aplicó OCR a {len(scanned)} página(s) escaneada(s)")
        else:
            needs_ocr = True
            warnings.append(
                f"{len(scanned)} página(s) parecen escaneadas (sin texto): requieren OCR. "
                "Instale Tesseract y reprocese el documento."
            )
    noise = _repeated_lines(page_texts)
    parts = []
    for i, ch in enumerate(chunks):
        md = ch.get("text", "")
        if noise:
            md = "\n".join(line for line in md.splitlines() if _norm_line(line) not in noise)
        parts.append(f"<!-- página {i + 1} -->\n\n{md.strip()}")
    tables: list[dict] = []
    if len(doc) <= 100:
        for i, page in enumerate(doc):
            try:
                found = page.find_tables()
            except Exception as exc:  # noqa: BLE001
                warnings.append(f"No se pudieron detectar tablas en la página {i + 1}: {type(exc).__name__}")
                continue
            for t_idx, table in enumerate(found.tables):
                rows = [[cell_text(c) for c in row] for row in table.extract()]
                if len(rows) < 2:
                    continue
                headers, body = rows[0], rows[1:]
                tables.append(
                    {
                        "name": f"Página {i + 1} · tabla {t_idx + 1}",
                        "page": i + 1,
                        "headers": headers,
                        "rows": body[:MAX_TABLE_ROWS_JSON],
                        "total_rows": len(body),
                    }
                )
    toc = doc.get_toc()
    meta = {k: v for k, v in (doc.metadata or {}).items() if v}
    if toc:
        meta["toc"] = [{"level": lvl, "title": title, "page": page} for lvl, title, page in toc[:300]]
    return Extraction(
        markdown="\n\n".join(parts).strip() + "\n",
        plain_text="\n\f\n".join(page_texts),
        tables=tables,
        pages=len(doc),
        warnings=warnings,
        needs_ocr=needs_ocr,
        meta=meta,
        noise_removed=sorted(noise)[:50],
    )


def _norm_line(line: str) -> str:
    return re.sub(r"\d+", "#", line.strip().lower())


def _repeated_lines(pages: list[str]) -> set[str]:
    """Encabezados/pies: líneas cortas en los 2 primeros o últimos renglones repetidas en ≥60 % de las páginas
    (con ≥3 páginas). Se consideran ruido de extracción y se excluyen del Markdown, nunca del original."""
    if len(pages) < 3:
        return set()
    counts: Counter[str] = Counter()
    for text in pages:
        lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
        edge = lines[:2] + lines[-2:]
        counts.update({_norm_line(ln) for ln in edge if len(ln) <= 60 and len(ln.split()) <= 8})
    threshold = max(3, int(len(pages) * 0.6))
    return {line for line, n in counts.items() if n >= threshold and line not in ("", "#")}


# --------------------------------------------------------------------------- Word
def extract_docx(content: bytes) -> Extraction:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        document = docx.Document(io.BytesIO(content))
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError(f"No se pudo abrir el documento Word: {exc}") from exc
    md: list[str] = []
    plain: list[str] = []
    tables: list[dict] = []
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(child, document)
            text = p.text.strip()
            if not text:
                continue
            plain.append(text)
            style = (p.style.name if p.style is not None else "") or ""
            level = _heading_level(style)
            if level:
                md.append(f"{'#' * level} {text}")
            elif _is_list(p, style):
                md.append(f"- {_runs_markdown(p)}")
            else:
                md.append(_runs_markdown(p))
        elif tag == "tbl":
            t = Table(child, document)
            rows = [[cell.text.strip() for cell in row.cells][:MAX_COLS] for row in t.rows]
            rows = [_dedupe_merged(r) for r in rows]
            if not rows:
                continue
            headers, body = rows[0], rows[1:]
            md.append(md_table(headers, body[:MAX_TABLE_ROWS_MD]))
            plain.append("\n".join("\t".join(r) for r in rows))
            tables.append(
                {
                    "name": f"Tabla {len(tables) + 1}",
                    "headers": headers,
                    "rows": body[:MAX_TABLE_ROWS_JSON],
                    "total_rows": len(body),
                }
            )
    noise = []
    for section in document.sections:
        for part in (section.header, section.footer):
            for p in part.paragraphs:
                if p.text.strip():
                    noise.append(p.text.strip())
    props = document.core_properties
    meta = {k: cell_text(getattr(props, k)) for k in ("title", "author", "created", "modified") if getattr(props, k)}
    return Extraction(
        markdown="\n\n".join(md).strip() + "\n",
        plain_text="\n".join(plain),
        tables=tables,
        meta=meta,
        noise_removed=sorted(set(noise))[:20],
    )


def _heading_level(style: str) -> int:
    s = style.lower()
    if s in ("title", "título"):
        return 1
    m = re.match(r"(heading|título|titulo|encabezado)\s*(\d)", s)
    return min(int(m.group(2)), 6) if m else 0


def _is_list(p, style: str) -> bool:
    if "list" in style.lower() or "lista" in style.lower():
        return True
    ppr = p._p.pPr  # noqa: SLF001
    return ppr is not None and ppr.numPr is not None


def _runs_markdown(p) -> str:
    out = []
    for run in p.runs:
        text = run.text
        if not text:
            continue
        if run.bold and text.strip():
            text = f"**{text.strip()}** "
        elif run.italic and text.strip():
            text = f"*{text.strip()}* "
        out.append(text)
    return "".join(out).strip() or p.text.strip()


def _dedupe_merged(row: list[str]) -> list[str]:
    """python-docx repite el texto de celdas combinadas horizontalmente; se deja solo la primera."""
    out = []
    for i, c in enumerate(row):
        out.append("" if i > 0 and c and c == row[i - 1] else c)
    return out


# --------------------------------------------------------------------------- Excel
def _format_value(value: Any, number_format: str) -> str:
    if value is None:
        return ""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        fmt = number_format or ""
        if "%" in fmt:
            return f"{value * 100:.4g}%"
        for sym in ("€", "$", "£", "¥"):
            if sym in fmt:
                return f"{value:,.2f} {sym}"
    return cell_text(value)


def _unit_of(fmt: str) -> str | None:
    if "%" in fmt:
        return "%"
    for sym in ("€", "$", "£", "¥"):
        if sym in fmt:
            return sym
    if any(t in fmt.lower() for t in ("yy", "dd", "mm/", "h:mm")):
        return "fecha"
    return None


def extract_xlsx(content: bytes) -> Extraction:
    import openpyxl

    try:
        wb_f = openpyxl.load_workbook(io.BytesIO(content), data_only=False)
        wb_v = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError(f"No se pudo abrir el libro Excel: {exc}") from exc
    md: list[str] = []
    plain: list[str] = []
    tables: list[dict] = []
    warnings: list[str] = []
    missing_cache = 0
    for ws in wb_f.worksheets:
        wsv = wb_v[ws.title]
        max_row = min(ws.max_row or 0, MAX_TABLE_ROWS_JSON + 20)
        max_col = min(ws.max_column or 0, MAX_COLS)
        if (ws.max_column or 0) > MAX_COLS:
            warnings.append(f"Hoja «{ws.title}»: solo se procesan las primeras {MAX_COLS} columnas")
        rows: list[list[str]] = []
        formulas: list[dict] = []
        units: dict[int, str] = {}
        for r_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=max_row, max_col=max_col), start=1):
            out_row = []
            for c_idx, cell in enumerate(row):
                raw = cell.value
                fmt = cell.number_format or ""
                if isinstance(raw, str) and raw.startswith("="):
                    cached = wsv.cell(row=r_idx, column=c_idx + 1).value
                    if cached is None:
                        missing_cache += 1
                    formulas.append({"cell": cell.coordinate, "formula": raw, "value": cell_text(cached)})
                    out_row.append(_format_value(cached, fmt) if cached is not None else raw)
                else:
                    out_row.append(_format_value(raw, fmt))
                unit = _unit_of(fmt)
                if unit and raw is not None:
                    units.setdefault(c_idx, unit)
            rows.append(out_row)
        while rows and not any(c.strip() for c in rows[-1]):
            rows.pop()
        if not rows:
            md.append(f"## Hoja: {ws.title}\n\n_(hoja vacía)_")
            continue
        headers, body = split_header(rows)
        col_units = {headers[i] if i < len(headers) and headers[i] else f"Col{i + 1}": u for i, u in units.items()}
        section = [f"## Hoja: {ws.title}", f"Dimensiones: {ws.dimensions} · {len(body)} filas de datos"]
        if col_units:
            section.append("Unidades por columna: " + ", ".join(f"{k}: {v}" for k, v in col_units.items()))
        section.append(md_table(headers, body[:MAX_TABLE_ROWS_MD]))
        if len(body) > MAX_TABLE_ROWS_MD:
            section.append(f"_… {len(body) - MAX_TABLE_ROWS_MD} filas más en la vista de tabla y en el original._")
        if formulas:
            section.append(f"### Fórmulas ({len(formulas)})")
            section.extend(
                f"- `{f['cell']}`: `{f['formula']}` → {f['value'] or '(sin valor calculado)'}"
                for f in formulas[:MAX_FORMULAS_MD]
            )
            if len(formulas) > MAX_FORMULAS_MD:
                section.append(f"- … {len(formulas) - MAX_FORMULAS_MD} fórmulas más")
        merged = [str(r) for r in ws.merged_cells.ranges]
        if merged:
            section.append("Celdas combinadas: " + ", ".join(merged[:40]))
        md.append("\n\n".join(section))
        plain.append(f"[{ws.title}]\n" + "\n".join("\t".join(r) for r in rows))
        tables.append(
            {
                "name": ws.title,
                "sheet": ws.title,
                "headers": headers,
                "rows": body[:MAX_TABLE_ROWS_JSON],
                "total_rows": len(body),
                "formulas": formulas[:2000],
                "merged": merged[:200],
                "units": col_units,
            }
        )
    if missing_cache:
        warnings.append(
            f"{missing_cache} fórmula(s) sin valor calculado en caché (el archivo no se recalculó en "
            "Excel); se muestra la fórmula"
        )
    names: list[str] = []
    with contextlib.suppress(AttributeError):
        names = [f"{name} = {dn.attr_text}" for name, dn in wb_f.defined_names.items()]
    meta = {"sheets": [ws.title for ws in wb_f.worksheets]}
    if names:
        meta["defined_names"] = names[:100]
        md.append("## Nombres definidos\n\n" + "\n".join(f"- `{n}`" for n in names[:100]))
    return Extraction(
        markdown="\n\n".join(md).strip() + "\n",
        plain_text="\n\n".join(plain),
        tables=tables,
        warnings=warnings,
        meta=meta,
    )


def extract_xls(content: bytes) -> Extraction:
    import xlrd

    try:
        book = xlrd.open_workbook(file_contents=content)
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError(f"No se pudo abrir el libro .xls: {exc}") from exc
    md, plain, tables = [], [], []
    for sh in book.sheets():
        rows = []
        for r in range(min(sh.nrows, MAX_TABLE_ROWS_JSON + 20)):
            row = []
            for c in range(min(sh.ncols, MAX_COLS)):
                cell = sh.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode).isoformat())
                else:
                    row.append(cell_text(cell.value) if cell.value != "" else "")
            rows.append(row)
        headers, body = split_header(rows)
        md.append(f"## Hoja: {sh.name}\n\n{md_table(headers, body[:MAX_TABLE_ROWS_MD])}")
        plain.append(f"[{sh.name}]\n" + "\n".join("\t".join(r) for r in rows))
        tables.append(
            {
                "name": sh.name,
                "sheet": sh.name,
                "headers": headers,
                "rows": body[:MAX_TABLE_ROWS_JSON],
                "total_rows": len(body),
            }
        )
    return Extraction(
        markdown="\n\n".join(md) + "\n",
        plain_text="\n\n".join(plain),
        tables=tables,
        warnings=["El formato .xls no conserva las fórmulas al leerlo; se muestran los valores"],
        meta={"sheets": book.sheet_names()},
    )


# --------------------------------------------------------------------------- CSV
def extract_csv(content: bytes, ext: str) -> Extraction:
    text, enc = decode_text(content)
    sample = text[:20000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|") if ext != "tsv" else csv.excel_tab
    except csv.Error:
        dialect = csv.excel
    rows = [[c.strip() for c in r][:MAX_COLS] for r in csv.reader(io.StringIO(text), dialect)]
    rows = [r for r in rows if any(r)]
    headers, body = split_header(rows) if rows else ([], [])
    if not headers and rows:
        headers, body = rows[0], rows[1:]
    md = [f"Tabla CSV · {len(body)} filas · codificación {enc}", md_table(headers, body[:MAX_TABLE_ROWS_MD])]
    if len(body) > MAX_TABLE_ROWS_MD:
        md.append(f"_… {len(body) - MAX_TABLE_ROWS_MD} filas más en la vista de tabla._")
    return Extraction(
        markdown="\n\n".join(md) + "\n",
        plain_text=text,
        tables=[{"name": "CSV", "headers": headers, "rows": body[:MAX_TABLE_ROWS_JSON], "total_rows": len(body)}],
        meta={"encoding": enc},
    )


# --------------------------------------------------------------------------- PowerPoint
def extract_pptx(content: bytes) -> Extraction:
    from pptx import Presentation

    try:
        prs = Presentation(io.BytesIO(content))
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError(f"No se pudo abrir la presentación: {exc}") from exc
    md, plain, tables = [], [], []
    for n, slide in enumerate(prs.slides, start=1):
        title = slide.shapes.title.text.strip() if slide.shapes.title is not None and slide.shapes.title.text else ""
        md.append(f"## Diapositiva {n}" + (f": {title}" if title else ""))
        plain.append(f"[Diapositiva {n}] {title}")
        for shape in slide.shapes:
            if shape == slide.shapes.title:
                continue
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = "".join(r.text for r in para.runs).strip()
                    if text:
                        md.append(("  " * para.level) + f"- {text}")
                        plain.append(text)
            if getattr(shape, "has_table", False) and shape.has_table:
                rows = [[cell.text.strip() for cell in row.cells] for row in shape.table.rows]
                if rows:
                    md.append(md_table(rows[0], rows[1:]))
                    tables.append(
                        {
                            "name": f"Diapositiva {n}",
                            "headers": rows[0],
                            "rows": rows[1:],
                            "total_rows": len(rows) - 1,
                            "slide": n,
                        }
                    )
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                md.append(f"> Notas: {notes}")
                plain.append(notes)
    return Extraction(
        markdown="\n\n".join(md) + "\n", plain_text="\n".join(plain), tables=tables, pages=len(prs.slides)
    )


# --------------------------------------------------------------------------- texto, datos y código
def extract_text(content: bytes, family: str, ext: str) -> Extraction:
    text, enc = decode_text(content)
    warnings = [f"Codificación detectada: {enc}"] if enc not in ("utf-8", "utf-8-sig") else []
    if family == "markdown":
        return Extraction(markdown=text if text.endswith("\n") else text + "\n", plain_text=text, warnings=warnings)
    if family == "text":
        return Extraction(markdown=text if text.endswith("\n") else text + "\n", plain_text=text, warnings=warnings)
    if family == "json":
        try:
            data = json.loads(text)
        except json.JSONDecodeError as exc:
            raise ExtractionError(f"JSON no válido: línea {exc.lineno}, columna {exc.colno}: {exc.msg}") from exc
        return _structured(text, data, "json", warnings)
    if family == "yaml":
        import yaml

        try:
            docs = list(yaml.safe_load_all(text))  # safe_load: nunca construye objetos arbitrarios
        except yaml.YAMLError as exc:
            raise ExtractionError(f"YAML no válido: {exc}") from exc
        return _structured(text, docs[0] if len(docs) == 1 else docs, "yaml", warnings)
    if family == "xml":
        from defusedxml import ElementTree

        try:
            root = ElementTree.fromstring(text.encode("utf-8"))
        except Exception as exc:  # noqa: BLE001 - defusedxml también bloquea entidades peligrosas
            raise ExtractionError(f"XML no válido o inseguro: {exc}") from exc
        outline = _xml_outline(root)
        body = text if len(text) < 200_000 else text[:200_000] + "\n<!-- truncado -->"
        md = f"## Estructura XML\n\n{outline}\n\n## Contenido\n\n```xml\n{body}\n```\n"
        return Extraction(markdown=md, plain_text=text, warnings=warnings, meta={"root": root.tag})
    lang = CODE_EXTS.get(ext, "")
    outline = _code_outline(text, lang)
    md = (f"## Esquema del código ({lang})\n\n{outline}\n\n" if outline else "") + f"```{lang}\n{text}\n```\n"
    return Extraction(
        markdown=md, plain_text=text, warnings=warnings, meta={"language": lang, "lines": text.count("\n") + 1}
    )


def _structured(text: str, data: Any, lang: str, warnings: list[str]) -> Extraction:
    outline = "\n".join(_data_outline(data))
    pretty = json.dumps(data, indent=2, ensure_ascii=False, default=str) if lang == "json" else text
    if len(pretty) > 200_000:
        pretty = pretty[:200_000] + "\n… (truncado)"
    md = f"## Estructura\n\n{outline}\n\n## Contenido\n\n```{lang}\n{pretty}\n```\n"
    return Extraction(markdown=md, plain_text=text, warnings=warnings)


def _data_outline(data: Any, prefix: str = "", depth: int = 0) -> list[str]:
    indent = "  " * depth
    if depth > 4:
        return []
    if isinstance(data, dict):
        out = []
        for k, v in list(data.items())[:60]:
            kind = type(v).__name__
            out.append(f"{indent}- `{k}` ({kind}{f', {len(v)} elementos' if isinstance(v, (list, dict)) else ''})")
            out.extend(_data_outline(v, f"{prefix}.{k}", depth + 1))
        return out
    if isinstance(data, list) and data:
        return [f"{indent}- [0] ({type(data[0]).__name__})", *_data_outline(data[0], prefix, depth + 1)]
    return []


def _xml_outline(root, depth: int = 0) -> str:
    lines = []

    def walk(el, d):
        if d > 4 or len(lines) > 200:
            return
        children = list(el)
        lines.append(f"{'  ' * d}- `{el.tag}`" + (f" ({len(children)} hijos)" if children else ""))
        seen = set()
        for ch in children:
            if ch.tag not in seen:
                seen.add(ch.tag)
                walk(ch, d + 1)

    walk(root, depth)
    return "\n".join(lines)


_DEF_PATTERNS = {
    "python": r"^\s*(?:async\s+)?(?:def|class)\s+(\w+)",
    "javascript": r"^\s*(?:export\s+)?(?:async\s+)?(?:function\s+(\w+)|class\s+(\w+)|const\s+(\w+)\s*=\s*\()",
    "typescript": r"^\s*(?:export\s+)?(?:async\s+)?(?:function\s+(\w+)|class\s+(\w+)|interface\s+(\w+)|type\s+(\w+))",
    "go": r"^\s*func\s+(?:\([^)]*\)\s*)?(\w+)",
    "rust": r"^\s*(?:pub\s+)?(?:fn|struct|enum|trait)\s+(\w+)",
    "java": r"^\s*(?:public|private|protected)?\s*(?:static\s+)?(?:class|interface|enum|[\w<>\[\]]+)\s+(\w+)\s*[({]",
    "csharp": (
        r"^\s*(?:public|private|protected|internal)?\s*(?:static\s+)?"
        r"(?:class|interface|[\w<>\[\]]+)\s+(\w+)\s*[({]"
    ),
    "php": r"^\s*(?:function|class)\s+(\w+)",
    "c": r"^\w[\w\s\*]+\s+(\w+)\s*\([^;]*$",
    "cpp": r"^\w[\w\s\*:<>]+\s+(\w+)\s*\([^;]*$",
}


def _code_outline(text: str, lang: str) -> str:
    pattern = _DEF_PATTERNS.get(lang if lang not in ("tsx", "jsx") else "typescript")
    if not pattern:
        return ""
    rx = re.compile(pattern)
    out = []
    for n, line in enumerate(text.splitlines(), start=1):
        m = rx.match(line)
        if m:
            name = next((g for g in m.groups() if g), None)
            if name:
                out.append(f"- `{name}` (línea {n})")
        if len(out) >= 200:
            break
    return "\n".join(out)


# --------------------------------------------------------------------------- despacho
def extract(content: bytes, family: str, ext: str) -> Extraction:
    if family == "pdf":
        return extract_pdf(content)
    if family == "docx":
        return extract_docx(content)
    if family == "xlsx":
        return extract_xlsx(content)
    if family == "xls":
        return extract_xls(content)
    if family == "csv":
        return extract_csv(content, ext)
    if family == "pptx":
        return extract_pptx(content)
    return extract_text(content, family, ext)
