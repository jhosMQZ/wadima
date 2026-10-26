"""Control horario: entrada y salida con la hora del servidor, jornadas incompletas, aprobación y totales.

Reglas:
- La hora de entrada y la de una salida normal salen del reloj del servidor de base de datos; el cliente nunca
  envía la hora de un registro normal.
- Una persona solo puede tener una jornada abierta (índice único parcial + bloqueo por usuario).
- Si alguien se desconecta sin registrar la salida no se inventa la hora: la jornada queda «pendiente de cierre».
  Opcionalmente (política `last_activity`) se cierra a la hora de la última conexión registrada, marcada como
  automática. Una salida declarada por el socio o fijada por un aprobador queda marcada con su tipo y motivo.
- El tiempo registrado es tiempo de conexión declarado, no tiempo efectivo de trabajo.
"""

from __future__ import annotations

import csv
import io
import uuid
from collections import defaultdict
from dataclasses import dataclass
from datetime import UTC, date, datetime, time, timedelta
from typing import Any
from zoneinfo import ZoneInfo

from sqlalchemy import or_, select, text
from sqlalchemy.orm import Session

from wadima import audit
from wadima import org as orgsvc
from wadima.activity import notify, touch_user_activity
from wadima.models import AreaMember, AuthSession, Organization, TimeEntry, TimeEvent, User, UserStatus
from wadima.realtime import emit
from wadima.security.auth import Principal, now_utc

CLOSE_KIND_LABEL = {
    "manual": "Registrada",
    "auto": "Automática (última conexión)",
    "declared": "Declarada por el socio",
    "admin": "Fijada por un aprobador",
}
APPROVAL_LABEL = {"pending": "Sin aprobar", "approved": "Aprobada", "rejected": "Rechazada"}
PENDING_LABEL = {"max_duration": "Supera la duración máxima", "disconnected": "Desconectado sin registrar salida"}


class TimeError(Exception):
    def __init__(self, status: int, message: str, entry: TimeEntry | None = None):
        super().__init__(message)
        self.status = status
        self.message = message
        self.entry = entry


# --------------------------------------------------------------------------- consultas básicas
def _lock_user(db: Session, user_id: uuid.UUID) -> None:
    """Serializa las operaciones horarias de una persona (dos clics o dos pestañas a la vez)."""
    db.execute(text("SELECT pg_advisory_xact_lock(hashtext(:k))"), {"k": f"time:{user_id}"})


def open_entry(db: Session, user_id: uuid.UUID) -> TimeEntry | None:
    return db.execute(
        select(TimeEntry).where(TimeEntry.user_id == user_id, TimeEntry.clock_out_at.is_(None))
    ).scalar_one_or_none()


def approvers(db: Session, org: Organization) -> set[uuid.UUID]:
    """Administradores y propietarios del área de Administración."""
    ids = set(db.execute(select(User.id).where(User.role == "admin", User.is_active)).scalars())
    area = orgsvc.template_area(db, "admin", org.id)
    if area is not None and area.enabled:
        ids |= set(
            db.execute(
                select(AreaMember.user_id)
                .join(User, User.id == AreaMember.user_id)
                .where(AreaMember.area_id == area.id, AreaMember.role == "owner", User.is_active)
            ).scalars()
        )
    return ids


def can_approve(db: Session, principal: Principal, org: Organization, entry: TimeEntry | None = None) -> bool:
    if entry is not None and entry.user_id == principal.id:
        return False  # nadie aprueba ni corrige su propia jornada
    return principal.id in approvers(db, org)


def can_view_user(db: Session, principal: Principal, org: Organization, user_id: uuid.UUID) -> bool:
    if user_id == principal.id or orgsvc.time_settings(org)["visibility"] == "all":
        return True
    return principal.id in approvers(db, org)


def _audience(db: Session, org: Organization, user_id: uuid.UUID) -> set[uuid.UUID] | None:
    if orgsvc.time_settings(org)["visibility"] == "all":
        return None
    return approvers(db, org) | {user_id}


def _event(
    db: Session,
    org: Organization,
    entry: TimeEntry,
    kind: str,
    *,
    at: datetime,
    actor_id: uuid.UUID | None,
    source: str,
    audit_id: int | None,
    detail: dict | None = None,
    note: str = "",
) -> TimeEvent:
    ev = TimeEvent(
        entry_id=entry.id,
        user_id=entry.user_id,
        actor_id=actor_id,
        kind=kind,
        at=at,
        source=source,
        detail=detail or {},
        audit_id=audit_id,
        note=note[:1000],
    )
    db.add(ev)
    db.flush()
    # Señal para la oficina (robot recolector) y los paneles: se publica solo tras el commit.
    audience = _audience(db, org, entry.user_id)
    emit(
        db,
        "time.recorded",
        audience=audience,
        everyone=audience is None,
        event_id=ev.id,
        kind=kind,
        user_id=str(entry.user_id),
        entry_id=str(entry.id),
    )
    return ev


def _db_now(db: Session) -> datetime:
    return db.execute(text("SELECT clock_timestamp()")).scalar_one()


# --------------------------------------------------------------------------- operaciones
def clock_in(
    db: Session, principal: Principal, org: Organization, *, ip: str | None, note: str = "", source: str = "web"
) -> tuple[TimeEntry, bool]:
    """Devuelve (jornada, creada). Si ya había una abierta no se crea otra y se deja constancia del intento."""
    _lock_user(db, principal.id)
    existing = open_entry(db, principal.id)
    if existing is not None:
        a = audit.record(
            db,
            "time.duplicate_blocked",
            actor_id=principal.id,
            target_type="time_entry",
            target_id=existing.id,
            ip=ip,
            data={"source": source},
        )
        _event(db, org, existing, "duplicate_blocked", at=a.ts, actor_id=principal.id, source=source, audit_id=a.id)
        return existing, False
    entry_id = uuid.uuid4()
    a = audit.record(
        db,
        "time.clock_in",
        actor_id=principal.id,
        target_type="time_entry",
        target_id=entry_id,
        ip=ip,
        data={"source": source},
    )
    entry = TimeEntry(
        id=entry_id,
        organization_id=org.id,
        user_id=principal.id,
        clock_in_at=a.ts,
        in_source=source,
        in_session_id=principal.session.id,
        in_ip=ip,
        in_note=note.strip()[:300],
        in_audit_id=a.id,
        approval_status="pending",
    )
    db.add(entry)
    db.flush()
    _event(db, org, entry, "clock_in", at=a.ts, actor_id=principal.id, source=source, audit_id=a.id)
    touch_user_activity(db, principal.id, "Registró su entrada")
    return entry, True


def clock_out(
    db: Session,
    principal: Principal,
    org: Organization,
    entry_id: uuid.UUID,
    *,
    ip: str | None,
    note: str = "",
    source: str = "web",
) -> TimeEntry:
    _lock_user(db, principal.id)
    entry = db.get(TimeEntry, entry_id)
    if entry is None or entry.user_id != principal.id:
        raise TimeError(404, "Jornada no encontrada")
    if entry.clock_out_at is not None:
        raise TimeError(409, "Esta jornada ya estaba cerrada", entry)
    a = audit.record(
        db,
        "time.clock_out",
        actor_id=principal.id,
        target_type="time_entry",
        target_id=entry.id,
        ip=ip,
        data={"source": source, "seconds": int((_db_now(db) - entry.clock_in_at).total_seconds())},
    )
    entry.clock_out_at = a.ts
    entry.close_kind = "manual"
    entry.out_source = source
    entry.out_session_id = principal.session.id
    entry.out_ip = ip
    entry.out_note = note.strip()[:300]
    entry.closed_by = principal.id
    entry.out_audit_id = a.id
    db.flush()
    _event(
        db,
        org,
        entry,
        "clock_out",
        at=a.ts,
        actor_id=principal.id,
        source=source,
        audit_id=a.id,
        detail={"seconds": duration_seconds(entry)},
    )
    touch_user_activity(db, principal.id, "Registró su salida")
    return entry


def _validate_close_time(db: Session, entry: TimeEntry, at: datetime) -> datetime:
    if at.tzinfo is None:
        raise TimeError(422, "Indique la hora con su zona horaria")
    at = at.astimezone(UTC)
    if at < entry.clock_in_at:
        raise TimeError(422, "La salida no puede ser anterior a la entrada")
    if at > _db_now(db):
        raise TimeError(422, "La salida no puede estar en el futuro")
    return at


def declare_exit(
    db: Session, principal: Principal, org: Organization, entry_id: uuid.UUID, at: datetime, reason: str, ip: str | None
) -> TimeEntry:
    """El socio olvidó registrar la salida e indica la hora real. Queda marcada como declarada."""
    _lock_user(db, principal.id)
    entry = db.get(TimeEntry, entry_id)
    if entry is None or entry.user_id != principal.id:
        raise TimeError(404, "Jornada no encontrada")
    if entry.clock_out_at is not None:
        raise TimeError(409, "Esta jornada ya estaba cerrada", entry)
    if len(reason.strip()) < 5:
        raise TimeError(422, "Explique brevemente por qué declara la hora (mínimo 5 caracteres)")
    at = _validate_close_time(db, entry, at)
    a = audit.record(
        db,
        "time.declared_close",
        actor_id=principal.id,
        target_type="time_entry",
        target_id=entry.id,
        ip=ip,
        data={"declared_at": at.isoformat(), "reason": reason.strip()[:300]},
    )
    entry.clock_out_at = at
    entry.close_kind = "declared"
    entry.close_reason = reason.strip()[:300]
    entry.out_source = "web"
    entry.out_session_id = principal.session.id
    entry.out_ip = ip
    entry.closed_by = principal.id
    entry.out_audit_id = a.id
    db.flush()
    _event(
        db,
        org,
        entry,
        "declared_close",
        at=a.ts,
        actor_id=principal.id,
        source="web",
        audit_id=a.id,
        detail={"declared_at": at.isoformat(), "seconds": duration_seconds(entry)},
        note=entry.close_reason,
    )
    notify(
        db,
        approvers(db, org) - {principal.id},
        "time.review",
        f"{principal.user.display_name} declaró la hora de salida de una jornada",
        f"Motivo: {entry.close_reason}",
        link="/time?tab=review",
    )
    return entry


def admin_close(
    db: Session, principal: Principal, org: Organization, entry_id: uuid.UUID, at: datetime, reason: str, ip: str | None
) -> TimeEntry:
    entry = db.get(TimeEntry, entry_id)
    if entry is None:
        raise TimeError(404, "Jornada no encontrada")
    if not can_approve(db, principal, org, entry):
        raise TimeError(403, "Solo un aprobador distinto del titular puede cerrar esta jornada")
    _lock_user(db, entry.user_id)
    db.refresh(entry)
    if entry.clock_out_at is not None:
        raise TimeError(409, "Esta jornada ya estaba cerrada", entry)
    if len(reason.strip()) < 5:
        raise TimeError(422, "Indique el motivo del cierre (mínimo 5 caracteres)")
    at = _validate_close_time(db, entry, at)
    a = audit.record(
        db,
        "time.admin_close",
        actor_id=principal.id,
        target_type="time_entry",
        target_id=entry.id,
        ip=ip,
        data={"closed_at": at.isoformat(), "reason": reason.strip()[:300], "user": str(entry.user_id)},
    )
    entry.clock_out_at = at
    entry.close_kind = "admin"
    entry.close_reason = reason.strip()[:300]
    entry.out_source = "admin"
    entry.out_ip = ip
    entry.closed_by = principal.id
    entry.out_audit_id = a.id
    db.flush()
    _event(
        db,
        org,
        entry,
        "admin_close",
        at=a.ts,
        actor_id=principal.id,
        source="admin",
        audit_id=a.id,
        detail={"closed_at": at.isoformat(), "seconds": duration_seconds(entry)},
        note=entry.close_reason,
    )
    notify(
        db,
        {entry.user_id},
        "time.closed",
        f"{principal.user.display_name} cerró una jornada suya que estaba abierta",
        f"Motivo: {entry.close_reason}",
        link="/time",
    )
    return entry


def review(
    db: Session, principal: Principal, org: Organization, entry_id: uuid.UUID, approve: bool, note: str, ip: str | None
) -> TimeEntry:
    entry = db.get(TimeEntry, entry_id)
    if entry is None:
        raise TimeError(404, "Jornada no encontrada")
    if not can_approve(db, principal, org, entry):
        raise TimeError(403, "Solo un aprobador distinto del titular puede revisar esta jornada")
    if entry.clock_out_at is None:
        raise TimeError(409, "Solo se aprueban jornadas cerradas", entry)
    if not approve and len(note.strip()) < 5:
        raise TimeError(422, "Indique el motivo del rechazo (mínimo 5 caracteres)")
    status = "approved" if approve else "rejected"
    a = audit.record(
        db,
        f"time.{'approve' if approve else 'reject'}",
        actor_id=principal.id,
        target_type="time_entry",
        target_id=entry.id,
        ip=ip,
        data={"previous": entry.approval_status, "note": note.strip()[:300], "user": str(entry.user_id)},
    )
    entry.approval_status = status
    entry.approved_by = principal.id
    entry.approved_at = a.ts
    entry.approval_note = note.strip()[:300]
    db.flush()
    _event(
        db,
        org,
        entry,
        "approve" if approve else "reject",
        at=a.ts,
        actor_id=principal.id,
        source="admin",
        audit_id=a.id,
        note=entry.approval_note,
    )
    notify(
        db,
        {entry.user_id},
        "time.reviewed",
        f"Jornada {'aprobada' if approve else 'rechazada'} por {principal.user.display_name}",
        entry.approval_note,
        link="/time",
    )
    return entry


def _has_live_session(db: Session, user_id: uuid.UUID, now: datetime) -> bool:
    return (
        db.execute(
            select(AuthSession.id)
            .where(AuthSession.user_id == user_id, AuthSession.revoked_at.is_(None), AuthSession.expires_at > now)
            .limit(1)
        ).first()
        is not None
    )


def sweep(db: Session, online_users: set[uuid.UUID]) -> dict[str, int]:
    """Marca jornadas incompletas y, si la política lo indica, las cierra a la hora de la última conexión."""
    result = {"flagged": 0, "auto_closed": 0, "cleared": 0}
    now = now_utc()
    for entry in db.execute(select(TimeEntry).where(TimeEntry.clock_out_at.is_(None))).scalars().all():
        org = orgsvc.get_org(db, entry.organization_id)
        policy = orgsvc.time_settings(org)
        st = db.get(UserStatus, entry.user_id)
        last_seen = st.last_seen_at if st and st.last_seen_at else entry.clock_in_at
        offline_for = now - max(last_seen, entry.clock_in_at)
        disconnected = entry.user_id not in online_users and (
            offline_for > timedelta(minutes=policy["disconnect_minutes"])
            or not _has_live_session(db, entry.user_id, now)
        )
        too_long = now - entry.clock_in_at > timedelta(hours=policy["max_open_hours"])
        if disconnected and policy["auto_close"] == "last_activity":
            _lock_user(db, entry.user_id)
            at = max(last_seen, entry.clock_in_at)
            a = audit.record(
                db,
                "time.auto_close",
                target_type="time_entry",
                target_id=entry.id,
                data={"closed_at": at.isoformat(), "rule": "last_activity", "user": str(entry.user_id)},
            )
            entry.clock_out_at = at
            entry.close_kind = "auto"
            entry.close_reason = "Cierre automático a la hora de la última conexión registrada"
            entry.out_source = "system"
            entry.out_audit_id = a.id
            db.flush()
            _event(
                db,
                org,
                entry,
                "auto_close",
                at=a.ts,
                actor_id=None,
                source="system",
                audit_id=a.id,
                detail={"closed_at": at.isoformat(), "last_seen_at": last_seen.isoformat()},
            )
            notify(
                db,
                {entry.user_id} | approvers(db, org),
                "time.auto_closed",
                "Jornada cerrada automáticamente por desconexión",
                "Se usó la hora de la última conexión registrada. Revise y apruebe o corrija.",
                link="/time",
            )
            result["auto_closed"] += 1
            continue
        reason = "max_duration" if too_long else "disconnected" if disconnected else None
        # «Supera la duración máxima» prevalece sobre «desconectado»: no se vuelve a marcar.
        if reason and entry.pending_reason not in (reason, "max_duration"):
            entry.pending_since = entry.pending_since or now
            entry.pending_reason = reason
            a = audit.record(
                db,
                "time.flag_pending",
                target_type="time_entry",
                target_id=entry.id,
                data={"reason": reason, "user": str(entry.user_id)},
            )
            _event(
                db,
                org,
                entry,
                "flag_pending",
                at=a.ts,
                actor_id=None,
                source="system",
                audit_id=a.id,
                detail={"reason": reason, "last_seen_at": last_seen.isoformat()},
            )
            notify(
                db,
                {entry.user_id},
                "time.pending",
                "Tiene una jornada pendiente de cierre",
                PENDING_LABEL[reason] + ". Registre la salida o declare la hora real.",
                link="/time",
            )
            result["flagged"] += 1
        elif reason is None and entry.pending_reason == "disconnected":
            # Volvió a conectarse: la marca deja de aplicar (el evento queda en el historial).
            entry.pending_since = None
            entry.pending_reason = None
            result["cleared"] += 1
    return result


# --------------------------------------------------------------------------- cálculo de duraciones y totales
def duration_seconds(entry: TimeEntry, now: datetime | None = None) -> int:
    end = entry.clock_out_at or now
    if end is None:
        return 0
    return max(0, int((end - entry.clock_in_at).total_seconds()))


def split_by_day(start: datetime, end: datetime, tz: ZoneInfo) -> dict[date, float]:
    """Reparte un intervalo entre los días locales que abarca (una jornada que cruza la medianoche suma en
    ambos días). Las restas se hacen en UTC para que los cambios de horario no alteren la duración."""
    out: dict[date, float] = defaultdict(float)
    cur = start.astimezone(tz)
    stop = end.astimezone(tz)
    while cur < stop:
        midnight = datetime.combine(cur.date() + timedelta(days=1), time(0), tzinfo=tz)
        seg_end = min(midnight, stop)
        out[cur.date()] += (seg_end.astimezone(UTC) - cur.astimezone(UTC)).total_seconds()
        cur = seg_end
    return out


def local_range(tz: ZoneInfo, start: date, end: date) -> tuple[datetime, datetime]:
    """Intervalo UTC [inicio del día `start`, inicio del día siguiente a `end`) en la zona de la organización."""
    a = datetime.combine(start, time(0), tzinfo=tz).astimezone(UTC)
    b = datetime.combine(end + timedelta(days=1), time(0), tzinfo=tz).astimezone(UTC)
    return a, b


def entries_between(db: Session, tz: ZoneInfo, start: date, end: date, user_ids: list[uuid.UUID] | None):
    a, b = local_range(tz, start, end)
    q = select(TimeEntry).where(
        TimeEntry.clock_in_at < b, or_(TimeEntry.clock_out_at.is_(None), TimeEntry.clock_out_at > a)
    )
    if user_ids is not None:
        q = q.where(TimeEntry.user_id.in_(user_ids))
    return db.execute(q.order_by(TimeEntry.clock_in_at)).scalars().all()


@dataclass
class _Bucket:
    registered: float = 0.0
    approved: float = 0.0
    in_progress: float = 0.0
    pending_approval: float = 0.0
    rejected: float = 0.0
    entries: int = 0

    def out(self) -> dict[str, int]:
        return {
            "registered_seconds": round(self.registered),
            "approved_seconds": round(self.approved),
            "in_progress_seconds": round(self.in_progress),
            "pending_approval_seconds": round(self.pending_approval),
            "rejected_seconds": round(self.rejected),
            "entries": self.entries,
        }


def summarize(entries, tz: ZoneInfo, start: date, end: date, now: datetime) -> dict[str, Any]:
    """Totales diarios, semanales (ISO, lunes a domingo) y mensuales dentro de [start, end].

    - registradas: jornadas cerradas no rechazadas; aprobadas: subconjunto aprobado.
    - en curso: jornadas abiertas hasta ahora (no se suman a las registradas).
    """
    days: dict[date, _Bucket] = defaultdict(_Bucket)
    for e in entries:
        end_at = e.clock_out_at or now
        for d, secs in split_by_day(e.clock_in_at, end_at, tz).items():
            if d < start or d > end:
                continue
            b = days[d]
            if e.clock_out_at is None:
                b.in_progress += secs
            elif e.approval_status == "rejected":
                b.rejected += secs
            else:
                b.registered += secs
                if e.approval_status == "approved":
                    b.approved += secs
                else:
                    b.pending_approval += secs
        first_day = e.clock_in_at.astimezone(tz).date()
        if start <= first_day <= end:
            days[first_day].entries += 1

    def roll(keyfn) -> dict[str, _Bucket]:
        acc: dict[str, _Bucket] = defaultdict(_Bucket)
        for d, b in days.items():
            t = acc[keyfn(d)]
            t.registered += b.registered
            t.approved += b.approved
            t.in_progress += b.in_progress
            t.pending_approval += b.pending_approval
            t.rejected += b.rejected
            t.entries += b.entries
        return acc

    weeks = roll(lambda d: f"{d.isocalendar().year}-W{d.isocalendar().week:02d}")
    months = roll(lambda d: f"{d.year}-{d.month:02d}")
    total = roll(lambda d: "total").get("total", _Bucket())
    return {
        "days": [{"date": d.isoformat(), **days[d].out()} for d in sorted(days)],
        "weeks": [{"week": k, **weeks[k].out()} for k in sorted(weeks)],
        "months": [{"month": k, **months[k].out()} for k in sorted(months)],
        "totals": total.out(),
    }


# --------------------------------------------------------------------------- serialización y exportación
def entry_status(e: TimeEntry) -> str:
    if e.clock_out_at is not None:
        return "closed"
    return "pending_close" if e.pending_since else "open"


def entry_out(e: TimeEntry, names: dict[uuid.UUID, str], tz: ZoneInfo, now: datetime) -> dict[str, Any]:
    return {
        "id": str(e.id),
        "user_id": str(e.user_id),
        "user_name": names.get(e.user_id, "—"),
        "local_date": e.clock_in_at.astimezone(tz).date().isoformat(),
        "clock_in_at": e.clock_in_at.isoformat(),
        "clock_out_at": e.clock_out_at.isoformat() if e.clock_out_at else None,
        "status": entry_status(e),
        "close_kind": e.close_kind,
        "duration_seconds": duration_seconds(e) if e.clock_out_at else None,
        "running_seconds": duration_seconds(e, now) if e.clock_out_at is None else None,
        "pending_since": e.pending_since.isoformat() if e.pending_since else None,
        "pending_reason": e.pending_reason,
        "approval_status": e.approval_status,
        "approved_by": str(e.approved_by) if e.approved_by else None,
        "approved_by_name": names.get(e.approved_by) if e.approved_by else None,
        "approved_at": e.approved_at.isoformat() if e.approved_at else None,
        "approval_note": e.approval_note,
        "in_note": e.in_note,
        "out_note": e.out_note,
        "close_reason": e.close_reason,
        "closed_by_name": names.get(e.closed_by) if e.closed_by else None,
        "in_source": e.in_source,
        "out_source": e.out_source,
    }


def event_out(ev: TimeEvent, names: dict[uuid.UUID, str]) -> dict[str, Any]:
    return {
        "id": ev.id,
        "entry_id": str(ev.entry_id) if ev.entry_id else None,
        "user_id": str(ev.user_id),
        "user_name": names.get(ev.user_id, "—"),
        "actor_name": names.get(ev.actor_id) if ev.actor_id else "Sistema",
        "kind": ev.kind,
        "at": ev.at.isoformat(),
        "source": ev.source,
        "detail": ev.detail,
        "note": ev.note,
        "audit_id": ev.audit_id,
    }


def hhmm(seconds: float | None) -> str:
    if seconds is None:
        return ""
    total = round(seconds / 60)
    return f"{total // 60}:{total % 60:02d}"


def _safe_cell(value: Any) -> Any:
    """Evita la inyección de fórmulas al abrir el CSV en una hoja de cálculo."""
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value


EXPORT_HEADERS = [
    "Socio",
    "Fecha",
    "Entrada",
    "Salida",
    "Duración (h:mm)",
    "Duración (horas)",
    "Estado",
    "Tipo de salida",
    "Aprobación",
    "Aprobado por",
    "Nota de entrada",
    "Nota de salida",
    "Motivo de cierre",
    "Origen entrada",
    "Origen salida",
    "ID",
]


def export_rows(entries, names: dict[uuid.UUID, str], tz: ZoneInfo, now: datetime) -> list[list[Any]]:
    rows = []
    for e in entries:
        secs = duration_seconds(e) if e.clock_out_at else None
        rows.append(
            [
                names.get(e.user_id, "—"),
                e.clock_in_at.astimezone(tz).date().isoformat(),
                e.clock_in_at.astimezone(tz).strftime("%Y-%m-%d %H:%M:%S"),
                e.clock_out_at.astimezone(tz).strftime("%Y-%m-%d %H:%M:%S") if e.clock_out_at else "",
                hhmm(secs),
                round(secs / 3600, 4) if secs is not None else "",
                {"open": "Abierta", "pending_close": "Pendiente de cierre", "closed": "Cerrada"}[entry_status(e)],
                CLOSE_KIND_LABEL.get(e.close_kind or "", ""),
                APPROVAL_LABEL[e.approval_status],
                names.get(e.approved_by, "") if e.approved_by else "",
                e.in_note,
                e.out_note,
                e.close_reason,
                e.in_source,
                e.out_source or "",
                str(e.id),
            ]
        )
    return rows


def preamble(org: Organization, start: date, end: date, who: str, generated: datetime) -> list[list[str]]:
    tz = orgsvc.zone(org)
    return [
        [f"{org.legal_name} · Control horario"],
        [f"Periodo: {start.isoformat()} a {end.isoformat()} · Personas: {who}"],
        [f"Zona horaria: {org.timezone} · Generado: {generated.astimezone(tz):%Y-%m-%d %H:%M} por {org.platform_name}"],
        ["El tiempo registrado refleja entradas y salidas declaradas; no equivale necesariamente a tiempo efectivo."],
        [],
    ]


def to_csv(head: list[list[str]], rows: list[list[Any]]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    for line in head:
        w.writerow(line)
    w.writerow(EXPORT_HEADERS)
    for r in rows:
        w.writerow([_safe_cell(v) for v in r])
    return ("﻿" + buf.getvalue()).encode("utf-8")  # BOM: Excel detecta UTF-8


def to_xlsx(org: Organization, head: list[list[str]], rows: list[list[Any]], summary: dict[str, Any]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Registros"
    for line in head:
        ws.append(line)
    ws.append(EXPORT_HEADERS)
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([_safe_cell(v) for v in r])
    ws.freeze_panes = ws.cell(row=len(head) + 2, column=1)
    for col, width in zip(
        "ABCDEFGHIJKLMNOP", (22, 12, 20, 20, 10, 10, 18, 26, 12, 18, 30, 30, 30, 10, 10, 38), strict=True
    ):
        ws.column_dimensions[col].width = width
    days = wb.create_sheet("Totales diarios")
    days.append(["Fecha", "Registradas (h:mm)", "Aprobadas (h:mm)", "En curso (h:mm)", "Rechazadas (h:mm)", "Jornadas"])
    for c in days[1]:
        c.font = Font(bold=True)
    for d in summary["days"]:
        days.append(
            [
                d["date"],
                hhmm(d["registered_seconds"]),
                hhmm(d["approved_seconds"]),
                hhmm(d["in_progress_seconds"]),
                hhmm(d["rejected_seconds"]),
                d["entries"],
            ]
        )
    agg = wb.create_sheet("Semanas y meses")
    agg.append(["Periodo", "Registradas (h:mm)", "Aprobadas (h:mm)", "Registradas (horas)", "Aprobadas (horas)"])
    for c in agg[1]:
        c.font = Font(bold=True)
    for key, items in (("week", summary["weeks"]), ("month", summary["months"])):
        for it in items:
            agg.append(
                [
                    it[key],
                    hhmm(it["registered_seconds"]),
                    hhmm(it["approved_seconds"]),
                    round(it["registered_seconds"] / 3600, 2),
                    round(it["approved_seconds"] / 3600, 2),
                ]
            )
    wb.properties.creator = org.platform_name
    wb.properties.title = f"Control horario · {org.legal_name}"
    wb.properties.subject = org.legal_name
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
