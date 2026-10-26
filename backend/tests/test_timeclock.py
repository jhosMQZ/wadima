from __future__ import annotations

import csv
import io
import threading
import uuid
from datetime import UTC, date, datetime, timedelta
from zoneinfo import ZoneInfo

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select, text
from sqlalchemy.exc import DBAPIError

from tests.conftest import UserClient
from wadima import timeclock as tc
from wadima.db import session_scope
from wadima.models import Notification, TimeEntry, TimeEvent, UserStatus

BOGOTA = ZoneInfo("America/Bogota")


def _kinds(user_id: str) -> list[str]:
    with session_scope() as db:
        return list(
            db.execute(select(TimeEvent.kind).where(TimeEvent.user_id == uuid.UUID(user_id)).order_by(TimeEvent.id))
            .scalars()
            .all()
        )


def _set_policy(admin, **policy):
    org = admin.get("/api/org").json()
    body = {
        k: org[k]
        for k in (
            "version",
            "platform_name",
            "legal_name",
            "short_mark",
            "description",
            "primary_color",
            "secondary_color",
            "timezone",
        )
    }
    body["time"] = {**org["time"], **policy}
    r = admin.put("/api/org", json=body)
    assert r.status_code == 200, r.text


def test_clock_in_uses_server_time_and_blocks_duplicates(partners):
    _, c2, _ = partners["clients"]
    assert c2.get("/api/time/status").json()["open"] is None
    before = datetime.now(UTC)
    # Una hora enviada por el cliente se ignora: manda el reloj del servidor.
    r = c2.post("/api/time/clock-in", json={"note": "Inicio", "clock_in_at": "2020-01-01T00:00:00Z"})
    assert r.status_code == 201, r.text
    entry = r.json()["entry"]
    t = datetime.fromisoformat(entry["clock_in_at"])
    assert abs((t - before).total_seconds()) < 10 and entry["status"] == "open" and entry["in_note"] == "Inicio"
    dup = c2.post("/api/time/clock-in", json={})
    assert dup.status_code == 409 and dup.json()["entry"]["id"] == entry["id"]
    with session_scope() as db:
        assert db.execute(select(func.count()).select_from(TimeEntry)).scalar() == 1
    assert _kinds(c2.id) == ["clock_in", "duplicate_blocked"]
    status = c2.get("/api/time/status").json()
    assert status["open"]["id"] == entry["id"] and status["timezone"] == "America/Bogota"


def test_concurrent_clock_ins_create_a_single_entry(app, partners):
    sessions = [UserClient(app, "socio2") for _ in range(5)]
    codes: list[int] = []
    barrier = threading.Barrier(len(sessions))

    def go(c):
        barrier.wait()
        codes.append(c.post("/api/time/clock-in", json={}).status_code)

    threads = [threading.Thread(target=go, args=(c,)) for c in sessions]
    for th in threads:
        th.start()
    for th in threads:
        th.join()
    assert sorted(codes) == [201, 409, 409, 409, 409]
    with session_scope() as db:
        assert db.execute(select(func.count()).select_from(TimeEntry)).scalar() == 1


def test_clock_out_closes_the_right_entry(app, partners):
    _, c2, c3 = partners["clients"]
    e2 = c2.post("/api/time/clock-in", json={}).json()["entry"]
    e3 = c3.post("/api/time/clock-in", json={}).json()["entry"]
    assert c2.post("/api/time/clock-out", json={"entry_id": str(uuid.uuid4())}).status_code == 404
    assert c2.post("/api/time/clock-out", json={"entry_id": e3["id"]}).status_code == 404  # jornada ajena
    r = c2.post("/api/time/clock-out", json={"entry_id": e2["id"], "note": "Fin"})
    assert r.status_code == 200
    out = r.json()["entry"]
    assert out["status"] == "closed" and out["close_kind"] == "manual" and out["duration_seconds"] >= 0
    assert out["out_note"] == "Fin" and out["approval_status"] == "pending"
    again = c2.post("/api/time/clock-out", json={"entry_id": e2["id"]})
    assert again.status_code == 409 and again.json()["entry"]["status"] == "closed"
    # La jornada de socio3 sigue abierta y sobrevive a una sesión nueva (los datos están en la base).
    assert UserClient(app, "socio3").get("/api/time/status").json()["open"]["id"] == e3["id"]
    assert _kinds(c2.id) == ["clock_in", "clock_out"]


def test_declared_exit_is_marked_and_validated(partners):
    c1, c2, _ = partners["clients"]
    entry = c2.post("/api/time/clock-in", json={}).json()["entry"]
    start = datetime.fromisoformat(entry["clock_in_at"])
    url = f"/api/time/entries/{entry['id']}/declare-exit"
    before = (start - timedelta(minutes=5)).isoformat()
    future = (datetime.now(UTC) + timedelta(hours=1)).isoformat()
    assert c2.post(url, json={"clock_out_at": before, "reason": "Olvidé salir"}).status_code == 422
    assert c2.post(url, json={"clock_out_at": future, "reason": "Olvidé salir"}).status_code == 422
    assert c2.post(url, json={"clock_out_at": start.isoformat(), "reason": "no"}).status_code == 422
    naive = start.replace(tzinfo=None).isoformat()
    assert c2.post(url, json={"clock_out_at": naive, "reason": "Olvidé salir"}).status_code == 422
    assert c1.post(url, json={"clock_out_at": start.isoformat(), "reason": "Olvidé salir"}).status_code == 404
    r = c2.post(url, json={"clock_out_at": start.isoformat(), "reason": "Olvidé registrar la salida"})
    assert r.status_code == 200
    out = r.json()["entry"]
    assert out["close_kind"] == "declared" and out["close_reason"] == "Olvidé registrar la salida"
    assert out["clock_out_at"] == out["clock_in_at"]
    with session_scope() as db:
        n = db.execute(select(Notification).where(Notification.user_id == uuid.UUID(c1.id))).scalar_one()
        assert n.kind == "time.review"
    review = c1.get("/api/time/incidents").json()["needs_review"]
    assert [e["id"] for e in review] == [entry["id"]]


def test_approval_rules(app, partners):
    c1, c2, c3 = partners["clients"]
    entry = c2.post("/api/time/clock-in", json={}).json()["entry"]
    url = f"/api/time/entries/{entry['id']}"
    assert c1.post(f"{url}/approve", json={}).status_code == 409  # todavía abierta
    c2.post("/api/time/clock-out", json={"entry_id": entry["id"]})
    assert c2.post(f"{url}/approve", json={}).status_code == 403  # nadie aprueba lo suyo
    assert c3.post(f"{url}/approve", json={}).status_code == 403  # un miembro no es aprobador
    assert c1.post(f"{url}/reject", json={"note": ""}).status_code == 422
    r = c1.post(f"{url}/approve", json={"note": "Correcto"})
    assert r.status_code == 200 and r.json()["entry"]["approval_status"] == "approved"
    # El propietario del área de Administración también aprueba (útil para revisar al administrador).
    tpl = c1.post("/api/org/templates/admin/enable").json()
    admin_area = next(t for t in tpl if t["key"] == "admin")["area_id"]
    assert c1.put(f"/api/areas/{admin_area}/members/{c3.id}", json={"role": "owner"}).status_code == 200
    own = c1.post("/api/time/clock-in", json={}).json()["entry"]
    c1.post("/api/time/clock-out", json={"entry_id": own["id"]})
    assert c1.post(f"/api/time/entries/{own['id']}/approve", json={}).status_code == 403
    r = c3.post(f"/api/time/entries/{own['id']}/reject", json={"note": "Duplicada por error"})
    assert r.status_code == 200 and r.json()["entry"]["approval_status"] == "rejected"
    assert c3.get("/api/time/status").json()["can_approve"] is True
    assert "approve" in _kinds(c2.id) and "reject" in _kinds(c1.id)


def test_records_are_immutable_for_the_application_role(partners):
    _, c2, _ = partners["clients"]
    entry = c2.post("/api/time/clock-in", json={}).json()["entry"]
    c2.post("/api/time/clock-out", json={"entry_id": entry["id"]})
    denied = "solo inserción|inmutables|ya fue registrada|no se pueden borrar|permission denied|permiso denegado"
    for sql in (
        "UPDATE time_entries SET clock_in_at = clock_in_at - interval '1 hour'",
        "UPDATE time_entries SET clock_out_at = clock_out_at + interval '3 hours'",
        "DELETE FROM time_entries",
        "UPDATE time_events SET kind = 'clock_in'",
        "DELETE FROM time_events",
    ):
        with pytest.raises(DBAPIError, match=denied), session_scope() as db:
            db.execute(text(sql))


def _open_entry(user_id: str, started: datetime, last_seen: datetime) -> str:
    with session_scope() as db:
        e = TimeEntry(user_id=uuid.UUID(user_id), clock_in_at=started, in_source="web", approval_status="pending")
        db.add(e)
        st = db.get(UserStatus, uuid.UUID(user_id)) or UserStatus(user_id=uuid.UUID(user_id))
        st.last_seen_at = last_seen
        db.add(st)
        db.flush()
        return str(e.id)


def test_incomplete_sessions_are_flagged_without_inventing_an_exit(partners):
    _, c2, c3 = partners["clients"]
    now = datetime.now(UTC)
    eid = _open_entry(c2.id, now - timedelta(hours=3), now - timedelta(hours=2))
    with session_scope() as db:
        assert tc.sweep(db, set()) == {"flagged": 1, "auto_closed": 0, "cleared": 0}
    with session_scope() as db:
        e = db.get(TimeEntry, uuid.UUID(eid))
        assert e.clock_out_at is None and e.pending_reason == "disconnected" and e.pending_since is not None
        assert tc.sweep(db, set())["flagged"] == 0  # no se vuelve a marcar
    st = c2.get("/api/time/status").json()["open"]
    assert st["status"] == "pending_close" and st["pending_reason"] == "disconnected"
    # Vuelve a conectarse: la marca deja de aplicar, pero el evento queda en el historial.
    with session_scope() as db:
        assert tc.sweep(db, {uuid.UUID(c2.id)})["cleared"] == 1
    assert _kinds(c2.id) == ["flag_pending"]
    # Una jornada que supera la duración máxima se marca aunque la persona siga conectada.
    _open_entry(c3.id, now - timedelta(hours=20), now)
    with session_scope() as db:
        tc.sweep(db, {uuid.UUID(c3.id), uuid.UUID(c2.id)})
    incidents = partners["clients"][0].get("/api/time/incidents").json()
    assert [e["user_id"] for e in incidents["open_too_long"]] == [c3.id]
    assert [e["pending_reason"] for e in incidents["pending_close"]] == ["max_duration"]


def test_auto_close_policy_uses_last_recorded_connection(partners):
    c1, c2, _ = partners["clients"]
    _set_policy(c1, auto_close="last_activity")
    now = datetime.now(UTC)
    last_seen = now - timedelta(hours=2)
    eid = _open_entry(c2.id, now - timedelta(hours=5), last_seen)
    with session_scope() as db:
        assert tc.sweep(db, set())["auto_closed"] == 1
    with session_scope() as db:
        e = db.get(TimeEntry, uuid.UUID(eid))
        assert e.clock_out_at == last_seen and e.close_kind == "auto" and e.approval_status == "pending"
    assert _kinds(c2.id) == ["auto_close"]
    review = c1.get("/api/time/incidents").json()["needs_review"]
    assert review[0]["close_kind"] == "auto"


def _entry(start: datetime, end: datetime | None, approval: str = "pending") -> TimeEntry:
    return TimeEntry(user_id=uuid.uuid4(), clock_in_at=start, clock_out_at=end, approval_status=approval)


def test_totals_are_mathematically_correct():
    b = lambda *a: datetime(*a, tzinfo=BOGOTA)  # noqa: E731
    entries = [
        _entry(b(2026, 10, 19, 8, 0), b(2026, 10, 19, 12, 30), "approved"),  # 4,5 h
        _entry(b(2026, 10, 19, 22, 0), b(2026, 10, 20, 2, 0)),  # cruza la medianoche: 2 h + 2 h
        _entry(b(2026, 10, 26, 9, 0), b(2026, 10, 26, 10, 0), "rejected"),  # 1 h rechazada
        _entry(b(2026, 10, 31, 23, 0), None),  # en curso: 1 h + 1 h hasta "ahora"
    ]
    now = b(2026, 11, 1, 1, 0)
    s = tc.summarize(entries, BOGOTA, date(2026, 10, 19), date(2026, 11, 1), now)
    days = {d["date"]: d for d in s["days"]}
    h = 3600
    assert days["2026-10-19"]["registered_seconds"] == 6.5 * h
    assert days["2026-10-19"]["approved_seconds"] == 4.5 * h
    assert days["2026-10-19"]["entries"] == 2
    assert days["2026-10-20"]["registered_seconds"] == 2 * h and days["2026-10-20"]["entries"] == 0
    assert days["2026-10-26"]["registered_seconds"] == 0 and days["2026-10-26"]["rejected_seconds"] == h
    assert days["2026-10-31"]["in_progress_seconds"] == h and days["2026-11-01"]["in_progress_seconds"] == h
    weeks = {w["week"]: w for w in s["weeks"]}
    assert weeks["2026-W43"]["registered_seconds"] == 8.5 * h and weeks["2026-W43"]["approved_seconds"] == 4.5 * h
    assert weeks["2026-W44"]["registered_seconds"] == 0 and weeks["2026-W44"]["in_progress_seconds"] == 2 * h
    months = {m["month"]: m for m in s["months"]}
    assert months["2026-10"]["registered_seconds"] == 8.5 * h and months["2026-10"]["in_progress_seconds"] == h
    assert months["2026-11"]["in_progress_seconds"] == h
    assert s["totals"] == {
        "registered_seconds": 8.5 * h,
        "approved_seconds": 4.5 * h,
        "in_progress_seconds": 2 * h,
        "pending_approval_seconds": 4 * h,
        "rejected_seconds": h,
        "entries": 4,
    }
    # Un periodo de un solo día cuenta solo la parte de la jornada que cae dentro.
    one = tc.summarize(entries, BOGOTA, date(2026, 10, 20), date(2026, 10, 20), now)
    assert one["totals"]["registered_seconds"] == 2 * h and one["totals"]["entries"] == 0


def test_durations_survive_daylight_saving_changes():
    madrid = ZoneInfo("Europe/Madrid")  # 25-10-2026: a las 03:00 vuelve a ser 02:00
    start = datetime(2026, 10, 25, 0, 0, tzinfo=madrid)
    end = datetime(2026, 10, 25, 6, 0, tzinfo=madrid)
    assert tc.split_by_day(start, end, madrid) == {date(2026, 10, 25): 7 * 3600}
    assert tc.duration_seconds(_entry(start.astimezone(UTC), end.astimezone(UTC))) == 7 * 3600


def test_export_csv_and_xlsx(partners):
    c1, c2, _ = partners["clients"]
    org = c1.get("/api/org").json()
    body = {k: org[k] for k in ("version", "platform_name", "short_mark", "description", "primary_color")}
    body |= {"secondary_color": org["secondary_color"], "timezone": org["timezone"], "legal_name": "VANTIA GROUP"}
    c1.put("/api/org", json=body)
    e = c2.post("/api/time/clock-in", json={"note": '=HYPERLINK("http://x")'}).json()["entry"]
    c2.post("/api/time/clock-out", json={"entry_id": e["id"]})
    r = c1.get("/api/time/export?format=csv")
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/csv")
    assert r.content.startswith("﻿".encode())
    rows = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig"))))
    assert rows[0] == ["VANTIA GROUP · Control horario"]
    assert any("no equivale necesariamente" in " ".join(row) for row in rows[:5])
    header = rows.index(tc.EXPORT_HEADERS)
    data = rows[header + 1]
    assert data[0] == "Socio2" and data[10].startswith("'=HYPERLINK")  # sin inyección de fórmulas
    x = c1.get("/api/time/export?format=xlsx")
    wb = load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames == ["Registros", "Totales diarios", "Semanas y meses"]
    values = [c.value for c in wb["Registros"][header + 2]]
    assert values[0] == "Socio2" and str(values[10]).startswith("'=")
    assert c1.get("/api/time/export?format=pdf").status_code == 422


def test_visibility_policy(partners):
    c1, c2, c3 = partners["clients"]
    c2.post("/api/time/clock-in", json={})
    assert len(c3.get("/api/time/entries").json()["entries"]) == 1  # por defecto los socios ven el historial
    assert c3.get(f"/api/time/entries?user_id={c2.id}").status_code == 200
    _set_policy(c1, visibility="admins")
    assert c3.get(f"/api/time/entries?user_id={c2.id}").status_code == 404
    assert c3.get("/api/time/entries").json()["entries"] == []
    assert len(c1.get("/api/time/entries").json()["entries"]) == 1
    assert c3.get("/api/time/summary").json()["people"] == []
    bad = c1.get("/api/time/entries?from=2026-10-10&to=2026-10-01")
    assert bad.status_code == 422


def test_robot_event_only_after_a_persisted_record(server, partners):
    from tests.test_realtime import _recv_until

    c1, c2, _ = partners["clients"]
    cookie = {"cookie": f"wadima_session={c1.http.cookies.get('wadima_session')}"}
    with server.websocket_connect("/ws", headers=cookie) as ws:
        assert ws.receive_json()["type"] == "hello"
        entry = c2.post("/api/time/clock-in", json={}).json()["entry"]
        msg = _recv_until(ws, "time.recorded", kind="clock_in")
        assert msg["user_id"] == c2.id and msg["entry_id"] == entry["id"]
        with session_scope() as db:  # el evento llega cuando el registro ya está guardado
            assert db.get(TimeEvent, msg["event_id"]).entry_id == uuid.UUID(entry["id"])
        # Una operación rechazada no genera señal para el robot.
        c2.post("/api/time/clock-out", json={"entry_id": str(uuid.uuid4())})
        c2.post("/api/time/clock-out", json={"entry_id": entry["id"]})
        assert _recv_until(ws, "time.recorded", kind="clock_out")["entry_id"] == entry["id"]
