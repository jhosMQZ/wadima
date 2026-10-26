"""Genera documentos de prueba reales (no sensibles) en memoria."""

from __future__ import annotations

import io


def make_pdf(pages: int = 3, scanned: bool = False, password: str | None = None) -> bytes:
    import pymupdf

    doc = pymupdf.open()
    for i in range(pages):
        page = doc.new_page()
        if scanned:
            pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 200, 100), False)
            pix.clear_with(200)
            page.insert_image(pymupdf.Rect(50, 50, 450, 250), pixmap=pix)
            continue
        page.insert_text((72, 40), "ACME Corp - Informe interno", fontsize=8)
        page.insert_text((72, 90), f"Seccion {i + 1}: Resultados del trimestre", fontsize=16)
        topics = ["ventas", "costes", "personal", "infraestructura", "clientes", "riesgos"]
        body = [
            f"Los ingresos del trimestre {i + 1} crecieron gracias al proyecto Orion.",
            f"El analisis de {topics[i % len(topics)]} muestra una evolucion positiva en la region {i + 1}.",
            f"La empresa invirtio {1000 * (i + 1)} euros en {topics[(i + 2) % len(topics)]} y seguridad.",
        ]
        for n, line in enumerate(body):
            page.insert_text((72, 130 + n * 16), line, fontsize=11)
        page.insert_text((72, 800), f"Pagina {i + 1} de {pages}", fontsize=8)
    kwargs = {}
    if password:
        kwargs = {"encryption": pymupdf.PDF_ENCRYPT_AES_256, "user_pw": password, "owner_pw": password}
    return doc.tobytes(**kwargs)


def make_docx() -> bytes:
    import docx

    d = docx.Document()
    d.add_heading("Plan de proyecto", level=1)
    d.add_paragraph("Este plan describe las fases del proyecto Atlas para el equipo de socios.")
    d.add_heading("Objetivos", level=2)
    d.add_paragraph("Reducir los costes operativos un 10%.", style="List Bullet")
    d.add_paragraph("Lanzar la plataforma en marzo.", style="List Bullet")
    table = d.add_table(rows=3, cols=2)
    for r, (a, b) in enumerate([("Fase", "Fecha"), ("Diseño", "2026-11-01"), ("Desarrollo", "2027-01-15")]):
        table.cell(r, 0).text = a
        table.cell(r, 1).text = b
    d.sections[0].footer.paragraphs[0].text = "Documento confidencial"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_xlsx() -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ventas"
    ws.append(["Producto", "Unidades", "Precio", "Total", "Margen"])
    ws.append(["Licencia", 10, 120.5, "=B2*C2", 0.25])
    ws.append(["Soporte", 4, 300, "=B3*C3", 0.4])
    ws.append(["Total", "=SUM(B2:B3)", None, "=SUM(D2:D3)", None])
    for row in (2, 3):
        ws.cell(row=row, column=3).number_format = "#,##0.00 €"
        ws.cell(row=row, column=5).number_format = "0%"
    ws.merge_cells("A6:C6")
    ws["A6"] = "Notas del trimestre"
    ws2 = wb.create_sheet("Resumen")
    ws2.append(["Indicador", "Valor"])
    ws2.append(["Clientes", 42])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_pptx() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Estrategia 2027"
    slide.placeholders[1].text_frame.text = "Expandir a dos mercados"
    slide.notes_slide.notes_text_frame.text = "Recordar mencionar el presupuesto"
    slide2 = prs.slides.add_slide(prs.slide_layouts[5])
    slide2.shapes.title.text = "Presupuesto"
    rows, cols = 2, 2
    table = slide2.shapes.add_table(rows, cols, Inches(1), Inches(2), Inches(6), Inches(1)).table
    table.cell(0, 0).text = "Partida"
    table.cell(0, 1).text = "Importe"
    table.cell(1, 0).text = "Marketing"
    table.cell(1, 1).text = "5000"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
